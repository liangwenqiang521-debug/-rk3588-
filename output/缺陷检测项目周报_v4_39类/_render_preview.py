"""
从 v4 PPTX 读取每页文字,生成预览拼图(供 IDE 预览面板查看).
"""
from pptx import Presentation
from pptx.util import Emu
from PIL import Image, ImageDraw, ImageFont
import os

PPTX = r"D:\workbuddy\缺陷检测项目\output\缺陷检测项目周报_v4_39类\缺陷检测项目周报_v4_39类.pptx"
OUT  = r"D:\workbuddy\缺陷检测项目\output\缺陷检测项目周报_v4_39类\预览内容长图.png"

font_candidates = [
    r"C:\Windows\Fonts\msyh.ttc",
    r"C:\Windows\Fonts\simhei.ttf",
    r"C:\Windows\Fonts\simsun.ttc",
]
FONT_PATH = None
for f in font_candidates:
    if os.path.exists(f):
        FONT_PATH = f
        break

prs = Presentation(PPTX)
sw_pt = prs.slide_width / 12700.0
sh_pt = prs.slide_height / 12700.0
ratio_w, ratio_h = 960, int(960 * sh_pt / sw_pt)
print(f"PPT: {len(prs.slides)} 页, 单页预览 {ratio_w}x{ratio_h}px")


def collect_text(slide):
    items = []
    for shp in slide.shapes:
        if shp.has_text_frame:
            for p in shp.text_frame.paragraphs:
                for r in p.runs:
                    txt = r.text.strip()
                    if not txt:
                        continue
                    bold = bool(r.font.bold)
                    try:
                        size = r.font.size.pt if r.font.size else 18
                    except Exception:
                        size = 18
                    color = "000000"
                    try:
                        if r.font.color and r.font.color.rgb:
                            color = str(r.font.color.rgb)
                    except Exception:
                        pass
                    top = Emu(shp.top or 0).pt
                    left = Emu(shp.left or 0).pt
                    items.append((top, left, size, bold, txt, color))
    items.sort(key=lambda x: (x[0], x[1]))
    return items


def find_box_fill(slide):
    best = None
    best_area = 0
    for shp in slide.shapes:
        if shp.has_text_frame:
            continue
        try:
            w = Emu(shp.width or 0).pt
            h = Emu(shp.height or 0).pt
        except Exception:
            continue
        if w < 400 or h < 200:
            continue
        try:
            if shp.fill.type == 1:
                col = str(shp.fill.fore_color.rgb)
            else:
                col = None
        except Exception:
            col = None
        if col and col != "000000":
            area = w * h
            if area > best_area:
                best_area = area
                best = col
    return best


def fnt(size, bold=False):
    if FONT_PATH:
        return ImageFont.truetype(FONT_PATH, size)
    return ImageFont.load_default()


# 15 页用 5x3 网格
cols, rows = 5, 3
margin = 24
total_w = cols * ratio_w + (cols + 1) * margin
total_h = rows * ratio_h + (rows + 1) * margin + 60
canvas = Image.new("RGB", (total_w, total_h), "#F5F7FA")
draw = ImageDraw.Draw(canvas)
title_fnt = ImageFont.truetype(FONT_PATH, 32) if FONT_PATH else ImageFont.load_default()
draw.text((margin, 12), "缺陷检测项目周报 v4  ·  39 类专题  ·  15 页预览拼图  ·  内容预览(非像素级)",
          fill="#1E3A5F", font=title_fnt)

for idx, slide in enumerate(prs.slides):
    r, c = idx // cols, idx % cols
    x0 = margin + c * (ratio_w + margin)
    y0 = margin + 60 + r * (ratio_h + margin)
    bg = find_box_fill(slide) or "#FFFFFF"
    page = Image.new("RGB", (ratio_w, ratio_h), bg)
    pd = ImageDraw.Draw(page)
    pd.rectangle([0, 0, ratio_w - 1, ratio_h - 1], outline="#9CB3C9", width=2)
    page_no_fnt = ImageFont.truetype(FONT_PATH, 18) if FONT_PATH else ImageFont.load_default()
    pd.text((10, 8), f"P{idx + 1:02d}", fill="#FFFFFF", font=page_no_fnt)
    is_dark = bg.upper() in ("0A1F44", "0D47A1", "0A1F44", "1E3A5F")
    for top, left_pt, size, bold, text, color in collect_text(slide):
        scale = 960 / 720
        sz = max(10, int(size * scale))
        x = int(left_pt * scale)
        y = int(top * scale)
        fcolor = "#" + color.upper() if color else "#000000"
        if is_dark:
            fcolor = "#FFFFFF"
        try:
            f = fnt(sz, bold)
        except Exception:
            f = page_no_fnt
        max_chars = max(10, int((ratio_w - x - 10) / max(sz * 0.6, 6)))
        if len(text) > max_chars:
            text = text[:max_chars] + "..."
        pd.text((x, y), text, fill=fcolor, font=f)
    canvas.paste(page, (x0, y0))

canvas.save(OUT, optimize=True)
print(f"OK: {OUT}, {os.path.getsize(OUT)/1024:.1f} KB, {canvas.size}")
